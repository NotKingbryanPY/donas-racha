#!/usr/bin/env python3
"""Convert a private Google Sheets XLSX download to the phase-5 export format.

Requires openpyxl. The result contains customer data and must stay outside Git.
"""

import argparse
import datetime as dt
import json
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl


SHEETS = (
    'Clientes', 'Registros', 'Insignias', 'ClienteInsignias',
    'TemporadasRacha', 'TiendaRecompensas', 'Canjes',
)


def convert(source: Path, destination: Path, spreadsheet_id: str, timezone_name: str):
    timezone = ZoneInfo(timezone_name)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)

    def encode(value):
        if isinstance(value, dt.datetime):
            if value.tzinfo is None:
                value = value.replace(tzinfo=timezone)
            return value.astimezone(dt.timezone.utc).isoformat(timespec='milliseconds').replace('+00:00', 'Z')
        if isinstance(value, dt.date):
            return encode(dt.datetime.combine(value, dt.time()))
        if isinstance(value, dt.time):
            return value.isoformat()
        return value

    payload = {
        'schemaVersion': 1,
        'source': 'GOOGLE_SHEETS',
        'spreadsheetId': spreadsheet_id,
        'spreadsheetName': source.stem,
        'spreadsheetTimeZone': timezone_name,
        'exportedAt': dt.datetime.now(dt.timezone.utc).isoformat(timespec='milliseconds').replace('+00:00', 'Z'),
        'sheets': {},
    }
    for name in SHEETS:
        if name not in workbook.sheetnames:
            payload['sheets'][name] = {'missing': True, 'headers': [], 'rows': []}
            continue
        records = workbook[name].iter_rows(values_only=True)
        headers = [encode(value) for value in next(records, ())]
        rows = []
        for record in records:
            if any(value is not None and value != '' for value in record):
                rows.append([encode(value) for value in record[:len(headers)]])
        payload['sheets'][name] = {'missing': False, 'headers': headers, 'rows': rows}

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(payload, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    return {name: len(payload['sheets'][name]['rows']) for name in SHEETS}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--spreadsheet-id', required=True)
    parser.add_argument('--time-zone', default='America/Bogota')
    args = parser.parse_args()
    counts = convert(args.input, args.output, args.spreadsheet_id, args.time_zone)
    print(json.dumps({'counts': counts}))
